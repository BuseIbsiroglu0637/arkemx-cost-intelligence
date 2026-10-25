"""Safe read-only .xlsx intake and review-first decimal/unit normalization."""
import io
import re
import unicodedata
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile, is_zipfile
from openpyxl import load_workbook

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 2000
MAX_COLS = 120

HEADER_ALIASES = {
    "description": {"imalat", "imalat aciklamasi", "aciklama", "is kalemi", "imalat cinsi", "yapilan is", "is tanimi"},
    "quantity": {"miktar", "metraj", "toplam miktar", "toplam metraj", "adet", "quantity"},
    "unit": {"birim", "olcu birimi", "birimi", "unit"},
}
UNIT_ALIASES = {
    "m2": "m²", "m²": "m²", "m3": "m³", "m³": "m³", "m": "m", "mt": "m",
    "metre": "m", "kg": "kg", "ton": "ton", "ad": "adet", "adet": "adet", "set": "set", "takim": "takım"
}


class ExcelValidationError(ValueError):
    pass


def normalized(text: str) -> str:
    text = str(text).strip().replace("İ", "I").replace("ı", "i")
    text = unicodedata.normalize("NFKD", text)
    return re.sub(r"\s+", " ", "".join(c for c in text if not unicodedata.combining(c)).lower())


def normalize_unit(value: Any) -> tuple[str | None, list[str]]:
    if value is None or str(value).strip() == "":
        return None, ["UNIT_MISSING"]
    value = normalized(str(value)).replace("²", "2").replace("³", "3")
    if value not in UNIT_ALIASES:
        return None, ["UNIT_UNRECOGNIZED"]
    return UNIT_ALIASES[value], []


def parse_quantity(value: Any) -> tuple[Decimal | None, list[str]]:
    if value is None or str(value).strip() == "":
        return None, ["QUANTITY_MISSING"]
    if isinstance(value, bool):
        return None, ["QUANTITY_INVALID"]
    if isinstance(value, (int, float, Decimal)):
        number = Decimal(str(value))
    else:
        raw = str(value).strip().replace("\u00a0", "").replace(" ", "")
        if not re.fullmatch(r"[+-]?[\d.,]+", raw):
            return None, ["QUANTITY_INVALID"]
        if "," in raw and "." in raw:
            decimal_separator = "," if raw.rfind(",") > raw.rfind(".") else "."
            groups_separator = "." if decimal_separator == "," else ","
            chunks = raw.split(decimal_separator)
            if len(chunks) != 2 or any(len(x) != 3 for x in chunks[0].lstrip("+-").split(groups_separator)[1:]):
                return None, ["QUANTITY_AMBIGUOUS"]
            raw = raw.replace(groups_separator, "").replace(decimal_separator, ".")
        elif raw.count(",") + raw.count(".") == 1:
            sep = "," if "," in raw else "."
            before, after = raw.split(sep)
            # '1.234' can mean 1234 or 1.234; require human input.
            if len(after) == 3 and len(before.lstrip("+-")) <= 3:
                return None, ["QUANTITY_AMBIGUOUS"]
            raw = raw.replace(",", ".")
        elif "," in raw or "." in raw:
            sep = "," if "," in raw else "."
            chunks = raw.lstrip("+-").split(sep)
            if any(len(x) != 3 for x in chunks[1:]):
                return None, ["QUANTITY_AMBIGUOUS"]
            raw = raw.replace(sep, "")
        try:
            number = Decimal(raw)
        except InvalidOperation:
            return None, ["QUANTITY_INVALID"]
    if not number.is_finite() or number < 0:
        return None, ["QUANTITY_INVALID"]
    if number.as_tuple().exponent < -4 or number >= Decimal("100000000000000"):
        return None, ["QUANTITY_OUT_OF_RANGE"]
    return number, []


def safe_name(name: str) -> str:
    return Path(name.replace("\\", "/")).name[:255] or "metraj.xlsx"


def open_sheet(content: bytes, sheet: str | None = None):
    if not content or len(content) > MAX_FILE_BYTES:
        raise ExcelValidationError("File must be between 1 byte and 10 MB")
    if not is_zipfile(io.BytesIO(content)):
        raise ExcelValidationError("Not a valid .xlsx ZIP container")
    # Guard against obvious decompression bombs before openpyxl parses the ZIP.
    # Real public uploads still need a dedicated scanning/limits service.
    with ZipFile(io.BytesIO(content)) as archive:
        files = archive.infolist()
        if len(files) > 1000 or sum(f.file_size for f in files) > 50 * 1024 * 1024:
            raise ExcelValidationError("Workbook exceeds decompressed size limits")
        if any(f.file_size > 30 * 1024 * 1024 or
               (f.compress_size and f.file_size / f.compress_size > 1000)
               for f in files):
            raise ExcelValidationError("Workbook contains oversized compressed entries")
    try:
        book = load_workbook(io.BytesIO(content), data_only=False, read_only=True, keep_links=False)
    except (ValueError, OSError, BadZipFile, KeyError, TypeError) as exc:
        raise ExcelValidationError("Invalid or unsupported .xlsx file") from exc
    if not book.sheetnames:
        book.close()
        raise ExcelValidationError("Workbook contains no sheets")
    if sheet is not None and sheet not in book.sheetnames:
        book.close()
        raise ExcelValidationError("Sheet not found")
    tab = book[sheet or book.sheetnames[0]]
    if tab.max_column and tab.max_column > MAX_COLS:
        book.close()
        raise ExcelValidationError(f"Maximum {MAX_COLS} columns allowed")
    return book, tab


def guess_map(headers: list[str]) -> dict[str, str | None]:
    result = {}
    for field, aliases in HEADER_ALIASES.items():
        result[field] = next((h for h in headers if normalized(h) in aliases), None)
    return result


def read_metraj(content: bytes, sheet: str | None = None,
                columns: dict[str, str] | None = None, header_row: int = 1) -> dict:
    if not 1 <= header_row <= 30:
        raise ExcelValidationError("Header row must be between 1 and 30")
    book, tab = open_sheet(content, sheet)
    try:
        rows_iter = tab.iter_rows(min_row=header_row, values_only=False)
        first = next(rows_iter, None)
        if first is None:
            raise ExcelValidationError("Workbook is empty")
        headers = [str(c.value).strip() if c.value is not None else f"Column {i+1}" for i, c in enumerate(first)]
        if len(set(headers)) != len(headers):
            raise ExcelValidationError("Duplicate headers: rename columns in Excel")
        mapping = guess_map(headers)
        if columns:
            mapping.update({k: v for k, v in columns.items() if k in mapping and v})
        if not mapping["description"] or not mapping["quantity"] or not mapping["unit"]:
            raise ExcelValidationError("Map description, quantity and unit columns")
        for val in mapping.values():
            if val not in headers:
                raise ExcelValidationError(f"Unknown column: {val}")
        if len(set(mapping.values())) < len(mapping):
            raise ExcelValidationError("Each field must map to a different column")
        indices = {k: headers.index(v) for k, v in mapping.items()}
        result = []
        for cells in rows_iter:
            if not any(c.value is not None and str(c.value).strip() for c in cells):
                continue
            if len(result) >= MAX_ROWS:
                raise ExcelValidationError(f"Maximum {MAX_ROWS} work items per file")
            desc_cell = cells[indices["description"]] if len(cells) > indices["description"] else None
            qty_cell = cells[indices["quantity"]] if len(cells) > indices["quantity"] else None
            unit_cell = cells[indices["unit"]] if len(cells) > indices["unit"] else None
            description = desc_cell.value if desc_cell else None
            qty = qty_cell.value if qty_cell else None
            unit = unit_cell.value if unit_cell else None
            issues = []
            if any(c is not None and c.data_type == "f" for c in (desc_cell, qty_cell, unit_cell)):
                issues.append("FORMULA_NOT_EVALUATED")
                if qty_cell is not None and qty_cell.data_type == "f":
                    qty = None
            amount, amount_issues = parse_quantity(qty)
            normalized_unit_value, unit_issues = normalize_unit(unit)
            issues.extend(amount_issues + unit_issues)
            if description is None or not str(description).strip():
                issues.append("DESCRIPTION_MISSING")
            result.append({
                "source_row": (desc_cell.row if desc_cell else len(result) + 2),
                "description": str(description or "").strip()[:10000],
                "raw_quantity": "" if qty is None else str(qty)[:200],
                "raw_unit": "" if unit is None else str(unit)[:50],
                "quantity": str(amount) if amount is not None and "FORMULA_NOT_EVALUATED" not in issues else None,
                "unit": normalized_unit_value if "FORMULA_NOT_EVALUATED" not in issues else None,
                "issues": sorted(set(issues)),
                "review_required": bool(issues)
            })
        if not result:
            raise ExcelValidationError("No work items found")
        return {"sheet": tab.title, "header_row": header_row, "headers": headers, "mapping": mapping,
                "row_count": len(result), "review_count": sum(r["review_required"] for r in result),
                "items": result}
    finally:
        book.close()
