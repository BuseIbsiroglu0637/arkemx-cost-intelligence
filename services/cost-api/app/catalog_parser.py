"""M3 catalog XLSX/CSV parser. No scraping and no published data bundled."""
import csv
import io
import re
from hashlib import sha256
from decimal import Decimal
from pathlib import Path
from openpyxl import load_workbook
from .excel import ExcelValidationError, normalized, normalize_unit, parse_quantity, open_sheet

MAX_CATALOG_BYTES = 5 * 1024 * 1024
MAX_CATALOG_ROWS = 5000
MAX_CATALOG_COLS = 24
FIELDS = ("code", "description", "unit", "price")
ALIASES = {
    "code": {"poz", "poz no", "poz numarasi", "poz kodu", "kod", "code"},
    "description": {"tanimi", "tanim", "poz tanimi", "imalat", "aciklama", "description"},
    "unit": {"birim", "olcu birimi", "unit", "birimi"},
    "price": {"birim fiyat", "birim fiyati", "fiyat", "price", "tl"},
}
CODE_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9./_\\-]{0,79}$")


class CatalogError(ValueError):
    pass


def validated_content(name: str, content: bytes) -> str:
    suffix = Path((name or "").lower()).suffix
    if suffix not in {".csv", ".xlsx"}:
        raise CatalogError("Only .xlsx or UTF-8 CSV accepted")
    if not content or len(content) > MAX_CATALOG_BYTES:
        raise CatalogError("Catalog file size must be 1 byte to 5 MB")
    return suffix


def headers_mapping(headers, columns=None):
    if not headers or len(headers) > MAX_CATALOG_COLS:
        raise CatalogError("Catalog requires 1–24 columns")
    if len({normalized(s) for s in headers}) != len(headers):
        raise CatalogError("Duplicate header names")
    result = {field: next((h for h in headers if normalized(h) in ALIASES[field]), None)
              for field in FIELDS}
    if columns:
        for field, value in columns.items():
            if field not in FIELDS:
                raise CatalogError("Unknown mapped field")
            result[field] = value or None
    if not all(result[field] for field in ("code", "description", "unit")):
        raise CatalogError("Map code, description and unit columns")
    chosen = [x for x in result.values() if x]
    if any(x not in headers for x in chosen) or len(set(chosen)) != len(chosen):
        raise CatalogError("Mapped columns must be distinct existing headers")
    return result


def _rows_from_xlsx(content: bytes):
    try:
        book, sheet = open_sheet(content)
    except ExcelValidationError as exc:
        raise CatalogError(str(exc)) from exc
    try:
        if sheet.max_column and sheet.max_column > MAX_CATALOG_COLS:
            raise CatalogError("Too many columns")
        iterator = sheet.iter_rows(values_only=False)
        head = next(iterator, None)
        if not head:
            raise CatalogError("Empty workbook")
        headers = [str(c.value or "").strip() for c in head]
        rows = []
        for source_row, line in enumerate(iterator, start=2):
            if not any(cell.value is not None for cell in line):
                continue
            if len(rows) >= MAX_CATALOG_ROWS:
                raise CatalogError("Maximum 5000 catalog rows")
            if any(cell.data_type == "f" for cell in line):
                raise CatalogError(f"Formula cell at source row {source_row}")
            rows.append([cell.value for cell in line])
        return headers, rows
    finally:
        book.close()


def _rows_from_csv(content: bytes):
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise CatalogError("CSV must be UTF-8") from exc
    try:
        sample = text[:4096]
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        reader = csv.reader(io.StringIO(text), dialect=dialect, strict=True)
        headers = [s.strip() for s in next(reader, [])]
        rows = []
        for row in reader:
            if not any(v.strip() for v in row):
                continue
            if len(row) != len(headers):
                raise CatalogError("CSV row has inconsistent column count")
            if len(rows) >= MAX_CATALOG_ROWS:
                raise CatalogError("Maximum 5000 catalog rows")
            rows.append(row)
    except (csv.Error, StopIteration) as exc:
        raise CatalogError("Invalid CSV delimiter or quoting") from exc
    return headers, rows


def parse_catalog(filename: str, content: bytes, columns=None):
    suffix = validated_content(filename, content)
    headers, rows = _rows_from_xlsx(content) if suffix == ".xlsx" else _rows_from_csv(content)
    mapping = headers_mapping(headers, columns)
    indices = {k: headers.index(v) if v else None for k, v in mapping.items()}
    unique_codes = set()
    entries = []
    for source_row, row in enumerate(rows, 2):
        code = str(row[indices["code"]] or "").strip()
        desc = str(row[indices["description"]] or "").strip()
        raw_unit = row[indices["unit"]]
        raw_price = row[indices["price"]] if indices["price"] is not None else None
        if not CODE_RE.fullmatch(code):
            raise CatalogError(f"Invalid position code on data row {source_row}")
        if code in unique_codes:
            raise CatalogError(f"Duplicate position code {code} on data row {source_row}")
        if not (3 <= len(desc) <= 2000) or desc.lstrip().startswith(("=", "+", "@")):
            raise CatalogError(f"Unsafe or missing description on data row {source_row}")
        unit, warnings = normalize_unit(raw_unit)
        if warnings or unit is None:
            raise CatalogError(f"Unknown unit on data row {source_row}")
        price = None
        if raw_price is not None and str(raw_price).strip():
            price, issues = parse_quantity(raw_price)
            if issues or price is None or price <= Decimal("0"):
                raise CatalogError(f"Ambiguous or invalid price on data row {source_row}")
        unique_codes.add(code)
        entries.append({"code": code, "description": desc, "unit": unit,
                        "price": str(price) if price is not None else None,
                        "source_row": source_row})
    if not entries:
        raise CatalogError("No catalog positions in file")
    return {"headers": headers, "mapping": mapping, "rows": entries,
            "sha256": sha256(content).hexdigest(), "row_count": len(entries)}


def inspect_catalog(filename: str, content: bytes):
    """Return headers before requiring a manual mapping."""
    suffix = validated_content(filename, content)
    headers, rows = _rows_from_xlsx(content) if suffix == ".xlsx" else _rows_from_csv(content)
    if not headers or len(headers) > MAX_CATALOG_COLS:
        raise CatalogError("Catalog must contain 1–24 headers")
    if len({normalized(s) for s in headers}) != len(headers):
        raise CatalogError("Duplicate header names")
    mapping = {field: next((h for h in headers if normalized(h) in ALIASES[field]), None)
               for field in FIELDS}
    return {"headers": headers, "suggested_mapping": mapping, "data_rows": len(rows)}
