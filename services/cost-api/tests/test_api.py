from io import BytesIO
from openpyxl import Workbook


def xlsx_bytes():
    w = Workbook()
    s = w.active
    s.append(["İmalat", "Metraj", "Ölçü Birimi"])
    s.append(["Seramik 60x60", "245,50", "m2"])
    s.append(["Gazbeton 10 cm", "1.234", "m2"])
    o = BytesIO()
    w.save(o)
    return o.getvalue()


def register(client, org, email):
    resp = client.post("/api/auth/register", json={"organization_name": org, "email": email,
                                                   "password": "very-secure-pilot-password"})
    assert resp.status_code == 201, resp.text
    return {"Authorization": "Bearer " + resp.json()["access_token"]}


def test_unauthorized(client):
    assert client.get("/api/projects").status_code == 401
    assert client.post("/api/projects", json={"name": "Nope"}).status_code == 401


def test_project_preview_import(client, auth):
    p = client.post("/api/projects", json={"name": "Kastamonu Pilot"}, headers=auth)
    assert p.status_code == 201, p.text
    pid = p.json()["id"]
    binary = xlsx_bytes()
    preview = client.post("/api/imports/preview", headers=auth,
                          files={"file": ("metraj.xlsx", binary)})
    assert preview.status_code == 200, preview.text
    assert preview.json()["row_count"] == 2
    assert preview.json()["review_count"] == 1
    saved = client.post(f"/api/projects/{pid}/imports", headers=auth,
                        files={"file": ("metraj.xlsx", binary)})
    assert saved.status_code == 201, saved.text
    assert len(saved.json()["sha256"]) == 64
    items = client.get(f"/api/projects/{pid}/items", headers=auth)
    assert items.status_code == 200
    assert len(items.json()) == 2
    assert items.json()[0]["quantity"] == "245.5000"
    assert items.json()[1]["review_required"] is True
    assert len(client.get(f"/api/projects/{pid}/imports", headers=auth).json()) == 1


def test_tenant_boundary(client):
    first = register(client, "Firm One", "a@example.com")
    second = register(client, "Firm Two", "b@example.com")
    project = client.post("/api/projects", json={"name": "Hidden Project"}, headers=first).json()["id"]
    assert client.get("/api/projects", headers=second).json() == []
    assert client.get(f"/api/projects/{project}/items", headers=second).status_code == 404
    assert client.get(f"/api/projects/{project}/imports", headers=second).status_code == 404
    assert client.post(f"/api/projects/{project}/imports", headers=second,
                       files={"file": ("metraj.xlsx", xlsx_bytes())}).status_code == 404


def test_source_permission_record_and_tenant(client):
    one = register(client, "Office A", "office-a@example.com")
    two = register(client, "Office B", "office-b@example.com")
    response = client.post("/api/catalog/sources", headers=one,
                           json={"institution": "SENTETIK TEST KATALOGU"})
    assert response.status_code == 201
    sid = response.json()["id"]
    assert response.json()["license_status"] == "pending"
    assert client.get("/api/catalog/sources", headers=two).json() == []
    assert client.post(f"/api/catalog/sources/{sid}/verify", headers=two,
                       json={"license_evidence": "Written consent signed", "permission_scope": "Pilot use"}).status_code == 404
    verified = client.post(f"/api/catalog/sources/{sid}/verify", headers=one,
                           json={"license_evidence": "Synthetic data produced for internal tests",
                                 "permission_scope": "Internal pilot without resale"})
    assert verified.status_code == 200
    assert verified.json()["license_status"] == "verified"


def test_invalid_upload(client, auth):
    project = client.post("/api/projects", json={"name": "Import test"}, headers=auth).json()["id"]
    bad = client.post(f"/api/projects/{project}/imports", headers=auth,
                      files={"file": ("bogus.xlsx", b"not a spreadsheet")})
    assert bad.status_code == 422
    wrong_ext = client.post("/api/imports/preview", headers=auth,
                            files={"file": ("bogus.xlsm", b"123")})
    assert wrong_ext.status_code == 415


def test_login_and_duplicate_email(client):
    first = register(client, "Alpha", "unique@example.com")
    ok = client.post("/api/auth/login", json={"email": "unique@example.com", "password": "very-secure-pilot-password"})
    assert ok.status_code == 200
    assert client.get("/api/me", headers=first).json()["role"] == "admin"
    assert client.post("/api/auth/login", json={"email": "unique@example.com", "password": "wrong"}).status_code == 401
    assert client.post("/api/auth/register", json={"organization_name": "Another", "email": "unique@example.com",
                                                   "password": "very-secure-pilot-password"}).status_code == 409


def test_inspect_unknown_headers_then_preview_with_mapping(client, auth):
    from openpyxl import Workbook
    w = Workbook()
    w.active.title = "İlk"
    w.active.append(["Karşılaştırma", "Değer", "Ölçü"])
    w.active.append(["Boya", 12, "m2"])
    second = w.create_sheet("İkinci")
    second.append(["Kalem", "Miktar", "Birim"])
    second.append(["Gazbeton", 5, "m3"])
    output = BytesIO(); w.save(output); binary = output.getvalue()
    inspect = client.post("/api/imports/inspect", headers=auth,
                          files={"file": ("metraj.xlsx", binary)}, data={"sheet": "İlk"})
    assert inspect.status_code == 200, inspect.text
    assert inspect.json()["sheets"] == ["İlk", "İkinci"]
    assert inspect.json()["mapping"]["description"] is None
    preview = client.post("/api/imports/preview", headers=auth,
                          files={"file": ("metraj.xlsx", binary)},
                          data={"sheet": "İlk", "description_column": "Karşılaştırma",
                                "quantity_column": "Değer", "unit_column": "Ölçü"})
    assert preview.status_code == 200, preview.text
    assert preview.json()["items"][0]["description"] == "Boya"
    assert preview.json()["items"][0]["quantity"] == "12"
